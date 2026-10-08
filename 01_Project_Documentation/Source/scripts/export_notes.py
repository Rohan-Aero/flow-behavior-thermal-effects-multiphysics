# -*- coding: utf-8 -*-
"""Section 10C-1 - export the speaker notes of the finished deck to Speaker_Notes/ (RE-ANALYSIS 2026).
Reads the notes back from the .pptx itself (so the export matches what PowerPoint shows) and writes
SPEAKER_NOTES.md with slide number, slide title, planned time and word count.
Usage: python export_notes.py <deck.pptx> <Speaker_Notes dir>"""
import re, sys, os
from pptx import Presentation

deck, out = sys.argv[1:3]
prs = Presentation(deck)
L = ["# Speaker Notes — Flow Behavior and Thermal Effects in Multiphysics Systems", "",
     "Final technical presentation, Eleation internship project (February–May 2025), %d slides. The notes are speaking guidance — what the audience" % len(prs.slides),
     "is looking at, the engineering point, the number that matters and the transition — not a script to read aloud.", ""]
tot_s = tot_w = 0
rows = []
for i, sl in enumerate(prs.slides, 1):
    # slide title = the text box placed at the top of the slide (y < 1 inch)
    tops = [sh for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip() and sh.top is not None and sh.top < 914400]
    title = sorted(tops, key=lambda sh: sh.top)[0].text_frame.text if tops else ""
    if i == 1:
        title = "Title"
    elif i == len(prs.slides):
        title = "Conclusion / Questions"
    note = sl.notes_slide.notes_text_frame.text.strip() if sl.has_notes_slide else ""
    m = re.match(r"\[~(\d+) s\]\s*", note)
    sec = int(m.group(1)) if m else 0
    words = len(note.split())
    tot_s += sec; tot_w += words
    rows.append((i, title, sec, words))
    L += ["## Slide %d — %s" % (i, title), "", "*Planned time: about %d s · %d words*" % (sec, words), "", note, ""]
summary = ["| Slide | Title | Planned time [s] | Words |", "|---|---|---|---|"] + ["| %d | %s | %d | %d |" % r for r in rows]
summary += ["| **Total** | | **%d (%.1f min)** | **%d** |" % (tot_s, tot_s / 60, tot_w), ""]
L[5:5] = summary
os.makedirs(out, exist_ok=True)
open(os.path.join(out, "SPEAKER_NOTES.md"), "w", encoding="utf-8").write("\n".join(L))
print("slides", len(rows), "with notes", sum(1 for r in rows if r[3] > 0), "total", tot_s, "s", tot_w, "words")
