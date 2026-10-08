# -*- coding: utf-8 -*-
"""Section 10C-1 - quality control of the final presentation (RE-ANALYSIS 2026).
Checks the finished .pptx itself:
  1. numbers  - every number on the slides, in the tables, chart titles/labels and speaker notes is traced to
                11_Final_Audit/MASTER_PROJECT_DATA.csv (rounding and unit-scale aware) or to a project document;
                the parametric chart series must equal the master values; slide-8 profile series must equal the project CSV
  2. wording  - forbidden / sensitive words with their sentence (safe, optimal, best, validated, factor of safety, adequate ...)
  3. content  - verbatim re-analysis note; modelling-level label next to the key values
  4. figures  - every picture is a file of Figures/ (SHA-256), which the manifest ties to a hash-verified project file
  5. layout   - slide count and sequential slide-number fields; every shape, picture and table inside the slide;
                minimum font sizes (slide text and chart text); estimated text fit of every text box (Carlito metrics)
  6. notes    - speaker notes on every slide, word counts and planned time
Usage: python qc_10C1.py <deck.pptx> <14_Presentation dir> <MASTER_PROJECT_DATA.csv> <out.json> <corpus dir> [...]"""
import os, re, sys, csv, json, bisect, hashlib, zipfile, math
from pptx import Presentation
from pptx.util import Emu
from PIL import ImageFont

DECK, PDIR, MCSV, OUT = sys.argv[1:5]
CORPUS = sys.argv[5:]
EMU = 914400
SW, SH = 13.333, 7.5
res = {"deck": os.path.basename(DECK), "deck_sha256": hashlib.sha256(open(DECK, "rb").read()).hexdigest().upper(),
       "master_sha256": hashlib.sha256(open(MCSV, "rb").read()).hexdigest().upper()}
prs = Presentation(DECK)
zf = zipfile.ZipFile(DECK)

# ------------------------------------------------------------------ collect content
slides = []
for i, sl in enumerate(prs.slides, 1):
    d = {"n": i, "texts": [], "tables": [], "pictures": [], "charts": [], "notes": "", "shapes": []}
    for sh in sl.shapes:
        box = (sh.left / EMU, sh.top / EMU, sh.width / EMU, sh.height / EMU) if sh.width is not None else None
        d["shapes"].append({"name": sh.name, "box": box, "type": str(sh.shape_type)})
        if sh.has_text_frame and sh.text_frame.text.strip() and sh.text_frame.text.strip() != str(i):   # skip the slide-number field
            d["texts"].append({"text": sh.text_frame.text, "box": box, "frame": sh.text_frame})
        if sh.has_table:
            d["tables"].append({"cells": [[c.text for c in r.cells] for r in sh.table.rows], "box": box, "table": sh.table})
        if sh.shape_type == 13:  # picture
            d["pictures"].append({"sha": hashlib.sha256(sh.image.blob).hexdigest().upper(), "box": box})
        if sh.has_chart:
            part = sh.chart.part.partname.lstrip("/")
            d["charts"].append({"xml": zf.read(part).decode("utf-8"), "box": box, "part": part})
    d["notes"] = sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else ""
    d["xml"] = zf.read(sl.part.partname.lstrip("/")).decode("utf-8")
    slides.append(d)


def chart_text(xml):
    t = re.findall(r"<a:t>([^<]*)</a:t>", xml)
    return [x for x in t if x.strip()]


# ------------------------------------------------------------------ 1. numbers
NUM = re.compile(r"(?<![A-Za-z_\d.])\d+(?:\.\d+)?(?:[eE][-+]?\d+)?")
rows = {}
with open(MCSV, encoding="utf-8") as fh:
    for r in csv.DictReader(l for l in fh if not l.startswith("#")):
        rows[r["ID"]] = r
MV = []
for r in rows.values():
    for col in ("Value", "Documented value"):
        s = (r.get(col) or "")
        s2 = s.replace(",", "") if re.fullmatch(r"[-0-9.,eE+ ]*", s) else s
        for tok in NUM.findall(s2):
            try:
                MV.append((abs(float(tok)), r["ID"]))
            except ValueError:
                pass
MV.sort()
MVv = [x for x, _ in MV]
SUP = str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹⁻−", "0123456789--")


def canon(s):
    s = s.lstrip("-+").replace(",", "")
    if "e" in s.lower():
        return "%.6g" % float(s)
    if "." in s:
        s = s.rstrip("0").rstrip(".")
    return s.lstrip("0") or "0"


def build_corpus():
    exact, rnd, n = {}, [], 0
    for root in CORPUS:
        for dp, dn, fn in os.walk(root):
            if any(x in dp for x in ("S10C1", "14_Presentation")):
                continue
            for f in fn:
                p = os.path.join(dp, f)
                ext = f.lower().rsplit(".", 1)[-1]
                try:
                    sz = os.path.getsize(p)
                except OSError:
                    continue
                if not (ext in ("md", "json", "txt", "tsv", "py", "tex") and sz < 3e6 or ext == "csv" and sz < 3e5):
                    continue
                t = open(p, encoding="utf-8", errors="ignore").read()
                n += 1
                t = re.sub(r"(?<=\d),(?=\d{3}\b)", "", t)
                t = re.sub(r"(\d+(?:\.\d+)?)\s*[×x·*]\s*10\s*\^?\{?([-−⁻]?[0-9⁰¹²³⁴⁵⁶⁷⁸⁹]+)\}?", lambda m: " %se%s " % (m.group(1), m.group(2).translate(SUP)), t)
                rel = os.path.relpath(p, root)
                for tok in NUM.findall(t):
                    c = canon(tok)
                    exact.setdefault(c, rel)
                    if ext in ("md", "csv", "json") and not re.search(r"monitor|residual|history|nodal|mode\d", rel, re.I):
                        try:
                            rnd.append((float(tok), rel))
                        except ValueError:
                            pass
    rnd.sort()
    return exact, rnd, n


CEX, CRD, NCF = build_corpus()
CRv = [x for x, _ in CRD]
res["corpus_files_scanned"] = NCF
SCALES = [(1, ""), (1e3, "x1e3"), (1e-3, "x1e-3"), (1e6, "x1e6"), (1e-6, "x1e-6"), (100, "x100"), (0.01, "x0.01")]


def decimals(tok):
    if "e" in tok:
        m, e = tok.lower().split("e")
        return (len(m.split(".")[1]) if "." in m else 0) - int(e)
    return len(tok.split(".")[1]) if "." in tok else 0


def trace(tok):
    x, d = float(tok), decimals(tok)
    for s, lab in SCALES:
        h = 0.5 * 10 ** (-d) * s * (1 + 1e-9)
        i = bisect.bisect_left(MVv, x * s - h)
        if i < len(MVv) and MVv[i] <= x * s + h:
            return "master", MV[i][1] + ((" " + lab) if lab else "")
    c = canon(tok)
    if c in CEX:
        return "document-exact", CEX[c]
    h = 0.5 * 10 ** (-d) * (1 + 1e-9)
    i, hits = bisect.bisect_left(CRv, x - h), []
    while i < len(CRv) and CRv[i] <= x + h:
        hits.append(CRD[i][1]); i += 1
    if hits:
        return "document-rounded", sorted(hits, key=lambda p: (not p.endswith(".md"), len(p)))[0]
    return "untraced", ""


def numbers_in(text):
    t = text.replace(" ", " ").replace("−", "-").replace("–", " – ")
    t = re.sub(r"\[~\d+ s\]", " ", t)                                                   # planned-time tags in the notes
    t = re.sub(r"(\d+(?:\.\d+)?)\s*×\s*10([⁻⁰¹²³⁴⁵⁶⁷⁸⁹]+)", lambda m: " %se%s " % (m.group(1), m.group(2).translate(SUP)), t)
    t = re.sub(r"\b(19|20)\d\d\b", " ", t)                                              # years
    t = re.sub(r"\b(LC|S|V|Q|T|P|C|R|K|F|A|D)\d+[A-Z]?\b", " ", t)                        # case / load-case / support / task identifiers
    t = re.sub(r"(Sections?|Chapters?|Slides?|Table|Figure|Section)\s+\d+[A-C]?", " ", t)
    t = re.sub(r"(?<=\d),(?=\d{3}(?!\d))", "", t)
    t = re.sub(r"[⁰¹²³⁴⁵⁶⁷⁸⁹₀₁₂₃₄₅₆₇₈₉]", " ", t)                                       # super/subscript digits (m², λ₁)
    out = []
    for m in re.finditer(r"(?<![A-Za-z_\d.])\d+(?:\.\d+)?(?:e-?\d+)?(?![A-Za-z\d])", t):
        out.append((m.group(0), re.sub(r"\s+", " ", t[max(0, m.start() - 50):m.end() + 30])))
    return out


num_rows = []
for d in slides:
    sources = [("slide", x["text"]) for x in d["texts"]] + [("table", c) for tb in d["tables"] for r in tb["cells"] for c in r] + \
              [("chart", c) for ch in d["charts"] for c in chart_text(ch["xml"])] + [("notes", d["notes"])]
    for where, text in sources:
        for tok, ctx in numbers_in(text):
            sig = len(re.sub(r"[^\d]", "", tok.split("e")[0]).lstrip("0"))
            cls, src = trace(tok if "e" in tok else canon(tok) if "." in tok or True else tok)
            if cls == "untraced" and sig <= 1:
                cls = "single-digit"
            num_rows.append({"slide": d["n"], "where": where, "number": tok, "sig_digits": sig, "class": cls, "traced_to": src, "context": ctx})
cnt, cnt3 = {}, {}
for r in num_rows:
    cnt[r["class"]] = cnt.get(r["class"], 0) + 1
    if r["sig_digits"] >= 3:
        cnt3[r["class"]] = cnt3.get(r["class"], 0) + 1
res["numbers"] = {"total": len(num_rows), "by_class": cnt, "by_class_3plus_sig": cnt3,
                  "untraced": [r for r in num_rows if r["class"] == "untraced"],
                  "document_rounded": [r for r in num_rows if r["class"] == "document-rounded"], "all": num_rows}

# chart series against the sources
data = json.load(open(os.path.join(PDIR, "Tables", "slide_data.json"), encoding="utf-8"))
series_check = []
for d in slides:
    for ch in d["charts"]:
        xml = ch["xml"]
        sers = re.findall(r"<c:ser>(.*?)</c:ser>", xml, re.S)
        for s in sers:
            name = re.search(r"<c:v>([^<]*)</c:v>", s).group(1)
            yv = re.search(r"<c:yVal>.*?<c:numCache>(.*?)</c:numCache>", s, re.S)
            ys = [float(v) for v in re.findall(r"<c:v>([^<]*)</c:v>", yv.group(1))] if yv else []
            series_check.append({"slide": d["n"], "series": name, "n": len(ys), "values": ys})
ok_param = True
for key, c in data["charts"].items():
    found = [s for s in series_check if s["slide"] == 15 and s["values"] == c["y"]]
    ok_param &= bool(found)
prof = data["profiles"]
ok_prof = all(any(s["values"] == prof[k] for s in series_check if s["slide"] == 8) for k in ("w_b", "w_max", "p_area", "Tb_mass", "Twi", "Two"))
res["chart_series"] = {"parametric_series_equal_master": ok_param, "profile_series_equal_project_csv": ok_prof, "series": len(series_check)}

# ------------------------------------------------------------------ 2. wording
WORDS = {"safe": r"\bsafe\w*|\bsafety\b", "optimal": r"\boptim\w*", "best": r"\bbest\b", "worst": r"\bworst\b",
         "validat": r"\bvalidat\w*", "factor of safety": r"factor of safety|safety factor", "adequa": r"\badequa\w*",
         "guarantee": r"\bguarant\w*", "certif": r"\bcertif\w*", "FSI": r"\bFSI\b", "collapse": r"\bcollapse\w*",
         "recovered": r"\brecover\w*", "proven": r"\bprove[dn]?\b", "revolutionary/cutting-edge": r"revolutionar|cutting[- ]edge|highly accurate"}
NEG = re.compile(r"\b(not|no|never|nor|without|cannot|unavailable|rather than|instead|would|none|neither|deliberately)\b|n't", re.I)
wording = []
for d in slides:
    for where, text in [("slide", x["text"]) for x in d["texts"]] + [("table", c) for tb in d["tables"] for r in tb["cells"] for c in r] + [("notes", d["notes"])]:
        flat = re.sub(r"\s+", " ", text)
        for k, pat in WORDS.items():
            for m in re.finditer(pat, flat, re.I):
                a = max(flat.rfind(". ", 0, m.start()), flat.rfind(": ", 0, m.start()), 0)
                b = flat.find(". ", m.end()); b = len(flat) if b < 0 else b
                sent = flat[a:b + 1].strip(". :")
                wording.append({"slide": d["n"], "where": where, "word": k, "negated": bool(NEG.search(sent)), "sentence": sent[:300]})
res["wording"] = {"occurrences": len(wording), "not_negated": [w for w in wording if not w["negated"]], "all": wording}

# ------------------------------------------------------------------ 3. content
NOTE = ("The original internship project files were unavailable at the time of re-analysis. The presented geometry, simulation setup and "
        "numerical results therefore represent re-analysed engineering work generated using ANSYS and associated calculations.")
alltext = {d["n"]: re.sub(r"\s+", " ", " ".join(x["text"] for x in d["texts"]).replace(" ", " ")) for d in slides}
res["re_analysis_note_verbatim_on_slide"] = [n for n, t in alltext.items() if NOTE in t]
LEVEL = {"562.58": r"medium|CFD|baseline|Maximum solid", "581.71": r"analytic|Section 2", "560.83": r"fine", "605.16": r"peak|LC2|static|von Mises",
         "657.2": r"analytic|restrained", "1.108": r"λ|lambda|S1"}
lev = []
for d in slides:
    slide_all = re.sub(r"\s+", " ", (" ".join(t["text"] for t in d["texts"]) + " " + " ".join(" ".join(r) for tb in d["tables"] for r in tb["cells"])).replace(" ", " "))
    for where, x in [("slide", t["text"]) for t in d["texts"]] + [("table", " | ".join(" | ".join(r) for r in tb["cells"])) for tb in d["tables"]] + [("notes", d["notes"])]:
        flat = re.sub(r"\s+", " ", x.replace(" ", " "))
        for val, pat in LEVEL.items():
            for m in re.finditer(re.escape(val), flat):
                win = flat[max(0, m.start() - 200):m.end() + 120]
                how = ("same text box / sentence" if re.search(pat, win, re.I) else
                       "same slide (card label, column or table header)" if where != "notes" and re.search(pat, slide_all, re.I) else "unlabelled")
                lev.append({"slide": d["n"], "where": where, "value": val, "label": how})
res["modelling_levels"] = {"occurrences": len(lev), "by_label": {k: sum(1 for x in lev if x["label"] == k) for k in {x["label"] for x in lev}},
                           "unlabelled": [x for x in lev if x["label"] == "unlabelled"], "all": lev}

# ------------------------------------------------------------------ 4. figures
man = json.load(open(os.path.join(PDIR, "Figures", "figure_manifest.json"), encoding="utf-8"))
figsha = {m["output_sha256"]: m for m in man}
pics = [(d["n"], p["sha"]) for d in slides for p in d["pictures"]]
res["figures"] = {"pictures_in_deck": len(pics), "distinct_images": len({s for _, s in pics}),
                  "all_from_Figures_folder": all(s in figsha for _, s in pics),
                  "all_project_verified": all(figsha[s]["matches_10A_manifest"] for _, s in pics if s in figsha),
                  "unknown_pictures": [n for n, s in pics if s not in figsha],
                  "per_slide": {n: [figsha[s]["slide_figure"] for m, s in pics if m == n and s in figsha] for n in range(1, len(slides) + 1)}}

# ------------------------------------------------------------------ 5. layout
res["layout"] = {"slides": len(slides),
                 "slide_number_field_on": [d["n"] for d in slides if 'type="slidenum"' in d["xml"]]}
oob = []
for d in slides:
    for sh in d["shapes"]:
        if sh["box"]:
            x, y, w, h = sh["box"]
            if x < -0.01 or y < -0.01 or x + w > SW + 0.01 or y + h > SH + 0.01:
                oob.append({"slide": d["n"], "shape": sh["name"], "box": [round(v, 2) for v in sh["box"]]})
res["layout"]["out_of_bounds"] = oob
# fonts on slides
minf = []
for d in slides:
    for x in d["texts"]:
        for p in x["frame"].paragraphs:
            for r in p.runs:
                if r.font.size is not None and r.font.size.pt < 10 and r.text.strip():
                    minf.append({"slide": d["n"], "pt": r.font.size.pt, "text": r.text[:60]})
res["layout"]["runs_below_10pt"] = minf
chart_sz = [int(v) / 100 for d in slides for ch in d["charts"] for v in re.findall(r'sz="(\d+)"', ch["xml"])]
res["layout"]["chart_font_min_pt"] = min(chart_sz) if chart_sz else None
# estimated text fit (Carlito = metric-compatible Calibri)
FD = "/usr/share/fonts/truetype/crosextra/"
_fonts = {}


def font(size, bold, italic):
    k = (round(size * 4) / 4, bold, italic)
    if k not in _fonts:
        f = "Carlito-" + ("BoldItalic" if bold and italic else "Bold" if bold else "Italic" if italic else "Regular") + ".ttf"
        _fonts[k] = ImageFont.truetype(FD + f, int(round(size * 100)))   # 100 px per pt -> width in pt*100
    return _fonts[k]


def frame_height(frame, width_in):
    """estimated rendered height [in] of a text frame of given width (single paragraph wrap, 1.2 line height)."""
    body = frame._txBody
    bp = body.find("{http://schemas.openxmlformats.org/drawingml/2006/main}bodyPr")
    li = int(bp.get("lIns", 91440)) / EMU; ri = int(bp.get("rIns", 91440)) / EMU
    ti = int(bp.get("tIns", 45720)) / EMU; bi = int(bp.get("bIns", 45720)) / EMU
    avail = (width_in - li - ri) * 72 * 100      # in pt*100
    total = ti * 72 + bi * 72
    for p in frame.paragraphs:
        runs = [r for r in p.runs] or []
        size = max([r.font.size.pt for r in runs if r.font.size is not None] or [18])
        words, line_w, lines = [], 0.0, 1
        for r in runs:
            fs = r.font.size.pt if r.font.size is not None else size
            f = font(fs, bool(r.font.bold), bool(r.font.italic))
            for tok in re.split(r"(\s+)", r.text):
                if not tok:
                    continue
                w = f.getlength(tok)
                if line_w + w > avail and tok.strip() and line_w > 0:
                    lines += 1; line_w = w if tok.strip() else 0
                else:
                    line_w += w
        sp = p.space_after.pt if p.space_after is not None else 0
        indent = 0.3 * 72 * 100 if p.level or ("buChar" in p._p.xml) else 0
        if indent:
            lines = max(lines, 1)
        total += lines * size * 1.2 + sp
    return total / 72


fit = []
for d in slides:
    for x in d["texts"]:
        if not x["box"]:
            continue
        need = frame_height(x["frame"], x["box"][2])
        if need > x["box"][3] + 0.12:
            fit.append({"slide": d["n"], "text": x["text"][:70].replace("\n", " / "), "box_h": round(x["box"][3], 2), "est_h": round(need, 2)})
res["layout"]["text_fit_flags"] = fit

# ------------------------------------------------------------------ 6. notes
nw = [(d["n"], len(d["notes"].split()), int(re.match(r"\[~(\d+) s\]", d["notes"]).group(1)) if re.match(r"\[~(\d+) s\]", d["notes"]) else 0) for d in slides]
res["notes"] = {"slides_with_notes": sum(1 for _, w, _ in nw if w > 0), "words": {n: w for n, w, _ in nw},
                "planned_seconds": {n: s for n, _, s in nw}, "planned_total_min": round(sum(s for *_, s in nw) / 60, 1),
                "all_between_30_and_90_s": all(30 <= s <= 90 for *_, s in nw)}

json.dump(res, open(OUT, "w", encoding="utf-8"), indent=1, ensure_ascii=False, default=str)
summary = {k: res[k] for k in ("re_analysis_note_verbatim_on_slide", "chart_series")}
summary.update({"numbers": res["numbers"]["by_class"], "numbers_3sig": res["numbers"]["by_class_3plus_sig"], "untraced": len(res["numbers"]["untraced"]),
                "wording_not_negated": len(res["wording"]["not_negated"]), "wording_total": res["wording"]["occurrences"],
                "levels_unlabelled": len(res["modelling_levels"]["unlabelled"]), "figures": {k: v for k, v in res["figures"].items() if k != "per_slide"},
                "slides": res["layout"]["slides"], "slidenum_fields": len(res["layout"]["slide_number_field_on"]),
                "out_of_bounds": len(oob), "runs_below_10pt": len(minf), "chart_font_min_pt": res["layout"]["chart_font_min_pt"],
                "text_fit_flags": len(fit), "notes": {k: v for k, v in res["notes"].items() if k in ("slides_with_notes", "planned_total_min", "all_between_30_and_90_s")}})
print(json.dumps(summary, indent=1, ensure_ascii=False))
